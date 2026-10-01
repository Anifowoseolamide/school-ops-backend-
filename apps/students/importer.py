"""Excel / CSV student import.

Flow
----
1. `preview`: read the file, detect which column is which, validate every row,
   and report what will be imported and which rows need attention. Nothing is
   saved except the import record and the uploaded file.
2. `commit`: create the valid rows (and their guardians) in one transaction.

Column headers are matched loosely ("Surname", "Last name", "LAST_NAME" all map
to `last_name`). The caller can override detection with an explicit mapping.
"""
import csv
import io
import re
from datetime import date, datetime, timedelta

from django.conf import settings
from django.db import transaction
from django.utils import timezone

from apps.schools.models import ClassLevel, ClassRoom

from .models import Gender, Guardian, ImportStatus, Student, StudentImport

FIELDS = [
    "admission_number",
    "first_name",
    "middle_name",
    "last_name",
    "full_name",
    "gender",
    "date_of_birth",
    "classroom",
    "admission_date",
    "address",
    "guardian_name",
    "guardian_relationship",
    "guardian_phone",
    "guardian_email",
]

ALIASES = {
    "admission_number": [
        "admission number", "admission no", "adm no", "admission num", "admno", "reg no",
        "registration number", "student id", "student number", "admission",
    ],
    "first_name": ["first name", "firstname", "given name", "forename"],
    "middle_name": ["middle name", "middlename", "other name", "other names", "othername"],
    "last_name": ["last name", "lastname", "surname", "family name"],
    "full_name": ["name", "full name", "fullname", "student name", "name of student", "student"],
    "gender": ["gender", "sex"],
    "date_of_birth": ["date of birth", "dob", "birth date", "birthday", "d o b"],
    "classroom": ["class", "class arm", "current class", "class name", "classroom", "arm"],
    "admission_date": ["admission date", "date admitted", "date of admission"],
    "address": ["address", "home address", "residential address"],
    "guardian_name": [
        "guardian name", "parent name", "guardian", "parent", "parent guardian name",
        "name of parent", "father name", "mother name", "parents name",
    ],
    "guardian_relationship": ["relationship", "guardian relationship", "relation"],
    "guardian_phone": [
        "guardian phone", "parent phone", "phone", "phone number", "parents phone", "guardian phone number",
        "parent phone number", "mobile", "gsm", "telephone", "contact",
    ],
    "guardian_email": ["guardian email", "parent email", "email", "email address", "parents email"],
}

REQUIRED = ["admission_number", "classroom"]


def normalise_header(header) -> str:
    text = str(header or "").strip().lower()
    text = re.sub(r"['’]", "", text)  # "Parent's Phone" -> "parents phone"
    text = re.sub(r"[\._\-/()#:]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def detect_mapping(headers: list[str], overrides: dict | None = None) -> dict:
    """Return {field: original header}."""
    mapping: dict[str, str] = {}
    lookup = {}
    for field, names in ALIASES.items():
        for name in names:
            lookup.setdefault(name, field)
    for header in headers:
        field = lookup.get(normalise_header(header))
        if field and field not in mapping:
            mapping[field] = header
    for header, field in (overrides or {}).items():
        if field in FIELDS and header in headers:
            mapping = {f: h for f, h in mapping.items() if h != header and f != field}
            mapping[field] = header
    return mapping


def read_table(uploaded_file, filename: str) -> tuple[list[str], list[dict]]:
    """Read a CSV or XLSX file into (headers, rows). Rows are dicts keyed by header."""
    name = filename.lower()
    uploaded_file.seek(0)
    if name.endswith(".csv"):
        raw = uploaded_file.read()
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        reader = csv.reader(io.StringIO(text))
        grid = [row for row in reader]
    elif name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        grid = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
    else:
        raise ValueError("Upload a .csv or .xlsx file.")

    # The header is the first row with at least two non-empty cells.
    header_index = next(
        (i for i, row in enumerate(grid) if sum(1 for c in row if c not in (None, "")) >= 2), None
    )
    if header_index is None:
        raise ValueError("The file appears to be empty.")
    headers = [str(h).strip() if h is not None else "" for h in grid[header_index]]
    rows = []
    for row in grid[header_index + 1:]:
        if all(c in (None, "") or str(c).strip() == "" for c in row):
            continue
        record = {}
        for i, header in enumerate(headers):
            if header:
                record[header] = row[i] if i < len(row) else None
        rows.append(record)
    return headers, rows


def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


_CLASS_RE = re.compile(r"^(JSS|JS|SSS|SS)([1-3])([A-Z0-9]+)$")


def parse_class(value: str) -> tuple[str, str] | None:
    compact = re.sub(r"[\s\-_/\.]+", "", value.upper())
    match = _CLASS_RE.match(compact)
    if not match:
        return None
    prefix, number, arm = match.groups()
    prefix = "JSS" if prefix in ("JSS", "JS") else "SS"
    level = f"{prefix}{number}"
    if level not in ClassLevel.values:
        return None
    return level, arm


def parse_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 1 < value < 80000:  # Excel serial date
        return date(1899, 12, 30) + timedelta(days=int(value))
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y", "%d/%m/%y", "%Y/%m/%d", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Unrecognised date '{text}'. Use DD/MM/YYYY.")


def parse_gender(value: str) -> str:
    v = value.strip().lower()
    if v in ("m", "male", "boy"):
        return Gender.MALE
    if v in ("f", "female", "girl"):
        return Gender.FEMALE
    if v == "":
        return ""
    raise ValueError(f"Unrecognised gender '{value}'. Use M or F.")


def normalise_phone(value: str) -> str:
    digits = re.sub(r"[^\d+]", "", value)
    if digits.startswith("+234"):
        digits = "0" + digits[4:]
    elif digits.startswith("234") and len(digits) == 13:
        digits = "0" + digits[3:]
    elif len(digits) == 10 and not digits.startswith("0"):
        digits = "0" + digits  # Excel often drops the leading zero
    return digits


def split_full_name(full: str) -> tuple[str, str, str]:
    """'OKAFOR Chinedu Emeka' -> last='Okafor', first='Chinedu', middle='Emeka'.

    Nigerian school registers usually put the surname first. If the name contains a
    comma ('Okafor, Chinedu') the part before the comma is the surname.
    """
    if "," in full:
        last, rest = [p.strip() for p in full.split(",", 1)]
        parts = rest.split()
        return (parts[0] if parts else ""), " ".join(parts[1:]), last
    parts = full.split()
    if len(parts) == 1:
        return parts[0], "", ""
    return parts[1], " ".join(parts[2:]), parts[0]


TITLES = {"mr", "mrs", "miss", "ms", "dr", "chief", "alhaji", "alhaja", "pastor", "engr", "prof", "barr", "rev"}


def split_guardian_name(name: str) -> tuple[str, str]:
    """'Mrs Ngozi Okafor' -> ('Mrs Ngozi', 'Okafor'); 'Mr Bello' -> ('Mr', 'Bello')."""
    parts = name.split()
    if not parts:
        return "", ""
    title = parts[0] if parts[0].lower().rstrip(".") in TITLES else ""
    rest = parts[1:] if title else parts
    if not rest:
        return title, ""
    if len(rest) == 1:
        return (title or rest[0]), (rest[0] if title else "")
    first = f"{title} {' '.join(rest[:-1])}".strip()
    return first, rest[-1]


def _title(value: str) -> str:
    return value.title() if value.isupper() or value.islower() else value


def validate_rows(school, rows: list[dict], mapping: dict) -> tuple[list[dict], list[dict]]:
    """Return (valid_records, errors). Each error: {row, errors, data}."""
    classes = {(c.level, c.arm): c for c in ClassRoom.objects.filter(school=school)}
    existing = set(Student.objects.filter(school=school).values_list("admission_number", flat=True))
    existing_lower = {a.lower() for a in existing}
    seen = set()
    valid, errors = [], []

    def get(row, field):
        header = mapping.get(field)
        return _clean(row.get(header)) if header else ""

    for index, row in enumerate(rows, start=2):  # row 1 is the header
        problems = []
        record = {"row": index}
        admission = get(row, "admission_number")
        first, middle, last = get(row, "first_name"), get(row, "middle_name"), get(row, "last_name")
        if not (first and last) and get(row, "full_name"):
            first, middle, last = split_full_name(get(row, "full_name"))
        record.update(first_name=_title(first), middle_name=_title(middle), last_name=_title(last))

        if not admission:
            problems.append("Admission number is missing.")
        elif admission.lower() in existing_lower:
            problems.append(f"Admission number {admission} already exists in the school.")
        elif admission.lower() in seen:
            problems.append(f"Admission number {admission} appears more than once in the file.")
        record["admission_number"] = admission

        if not record["first_name"] or not record["last_name"]:
            problems.append("First name and last name (or a full name) are required.")

        class_value = get(row, "classroom")
        parsed = parse_class(class_value) if class_value else None
        if not class_value:
            problems.append("Class is missing.")
        elif parsed is None:
            problems.append(f"Could not read class '{class_value}'. Use a format like 'JSS 2A' or 'SS 1B'.")
        elif parsed not in classes:
            problems.append(f"Class '{class_value}' does not exist yet. Create it before importing.")
        else:
            record["classroom_id"] = classes[parsed].id
            record["classroom_name"] = classes[parsed].name

        try:
            record["gender"] = parse_gender(get(row, "gender"))
        except ValueError as exc:
            problems.append(str(exc))
        header = mapping.get("date_of_birth")
        try:
            record["date_of_birth"] = parse_date(row.get(header)) if header else None
        except ValueError as exc:
            problems.append(f"Date of birth: {exc}")
        header = mapping.get("admission_date")
        try:
            record["admission_date"] = parse_date(row.get(header)) if header else None
        except ValueError as exc:
            problems.append(f"Admission date: {exc}")

        record["address"] = get(row, "address")
        phone = normalise_phone(get(row, "guardian_phone"))
        if phone and not (10 <= len(phone.lstrip("+")) <= 15):
            problems.append(f"Guardian phone '{get(row, 'guardian_phone')}' does not look like a phone number.")
        record["guardian_phone"] = phone
        record["guardian_name"] = _title(get(row, "guardian_name"))
        record["guardian_relationship"] = get(row, "guardian_relationship")
        record["guardian_email"] = get(row, "guardian_email").lower()

        if problems:
            errors.append({"row": index, "errors": problems, "data": {k: _clean(v) for k, v in row.items()}})
        else:
            seen.add(admission.lower())
            valid.append(record)
    return valid, errors


def _jsonable_record(record: dict) -> dict:
    return {k: (v.isoformat() if isinstance(v, date) else v) for k, v in record.items()}


def preview_import(school, user, uploaded_file, mapping_overrides=None) -> tuple[StudentImport, dict]:
    filename = uploaded_file.name
    headers, rows = read_table(uploaded_file, filename)
    max_rows = settings.IMPORT_MAX_ROWS
    if len(rows) > max_rows:
        raise ValueError(f"The file has {len(rows)} rows. The maximum per import is {max_rows}.")
    mapping = detect_mapping(headers, mapping_overrides)
    missing = [f for f in REQUIRED if f not in mapping]
    if "first_name" not in mapping and "full_name" not in mapping:
        missing.append("first_name/full_name")
    valid, errors = validate_rows(school, rows, mapping) if not missing else ([], [])
    uploaded_file.seek(0)
    job = StudentImport.objects.create(
        school=school,
        uploaded_by=user,
        file=uploaded_file,
        original_filename=filename[:255],
        column_mapping=mapping,
        total_rows=len(rows),
        valid_rows=len(valid),
        error_rows=errors[:1000],
        status=ImportStatus.PREVIEWED if not missing else ImportStatus.FAILED,
    )
    unmapped = [h for h in headers if h and h not in mapping.values()]
    report = {
        "import_id": job.id,
        "status": job.status,
        "filename": filename,
        "column_mapping": mapping,
        "unmapped_columns": unmapped,
        "missing_required_columns": missing,
        "total_rows": len(rows),
        "valid_rows": len(valid),
        "error_count": len(errors),
        "errors": errors[:200],
        "sample": [_jsonable_record(r) for r in valid[:20]],
    }
    return job, report


@transaction.atomic
def commit_import(job: StudentImport) -> StudentImport:
    """Create the valid rows of a previewed import. Rows with problems are skipped."""
    job.file.open("rb")
    try:
        headers, rows = read_table(job.file, job.original_filename)
    finally:
        job.file.close()
    valid, errors = validate_rows(job.school, rows, job.column_mapping)
    guardians_by_phone = {
        g.phone: g for g in Guardian.objects.filter(school=job.school).exclude(phone="").order_by("id")
    }
    created = 0
    for record in valid:
        student = Student.objects.create(
            school=job.school,
            admission_number=record["admission_number"],
            first_name=record["first_name"],
            middle_name=record["middle_name"],
            last_name=record["last_name"],
            gender=record.get("gender", ""),
            date_of_birth=record.get("date_of_birth"),
            admission_date=record.get("admission_date"),
            classroom_id=record["classroom_id"],
            address=record.get("address", ""),
        )
        phone = record.get("guardian_phone")
        if phone or record.get("guardian_name"):
            guardian = guardians_by_phone.get(phone) if phone else None
            if guardian is None:
                first, last = split_guardian_name(record.get("guardian_name") or "")
                if not first:
                    first, last = "Parent of", student.first_name
                guardian = Guardian.objects.create(
                    school=job.school,
                    first_name=first,
                    last_name=last,
                    relationship=record.get("guardian_relationship", ""),
                    phone=phone or "",
                    email=record.get("guardian_email", ""),
                )
                if phone:
                    guardians_by_phone[phone] = guardian
            student.guardians.add(guardian)
        created += 1
    job.created_count = created
    job.valid_rows = len(valid)
    job.error_rows = errors[:1000]
    job.status = ImportStatus.COMPLETED
    job.completed_at = timezone.now()
    job.save(update_fields=["created_count", "valid_rows", "error_rows", "status", "completed_at", "updated_at"])
    return job
